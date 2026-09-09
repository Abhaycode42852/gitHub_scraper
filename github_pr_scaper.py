import os
import sys
import time
from datetime import datetime

import requests
from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

GITHUB_API = "https://api.github.com"

load_dotenv()

GITHUB_TOKEN = os.getenv("GITHUB_TOKEN")
GITHUB_ORG = os.getenv("GITHUB_ORG")
GITHUB_USERNAME = os.getenv("GITHUB_USERNAME")


def github_request(url, params=None):
    headers = {
        "Authorization": f"Bearer {GITHUB_TOKEN}",
        "Accept": "application/vnd.github+json",
        "X-GitHub-Api-Version": "2022-11-28",
    }

    response = requests.get(
        url,
        headers=headers,
        params=params,
        timeout=30,
    )

    if response.status_code == 403:
        remaining = response.headers.get("X-RateLimit-Remaining")

        if remaining == "0":
            reset_time = int(
                response.headers.get("X-RateLimit-Reset", 0)
            )

            wait_seconds = max(
                reset_time - int(time.time()),
                1,
            )

            print(
                f"Rate limit reached. Waiting {wait_seconds} seconds..."
            )

            time.sleep(wait_seconds)

            return github_request(url, params)

    response.raise_for_status()

    return response.json()


def parse_datetime(value):
    if not value:
        return None

    return datetime.fromisoformat(
        value.replace("Z", "+00:00")
    )


def get_all_repositories(org):
    repositories = []
    page = 1

    print(f"\nSearching repositories in {org}...")

    while True:
        url = f"{GITHUB_API}/orgs/{org}/repos"

        params = {
            "per_page": 100,
            "page": page,
            "type": "all",
        }

        data = github_request(url, params)

        if not data:
            break

        repositories.extend(data)

        print(
            f"Found {len(repositories)} repositories so far..."
        )

        if len(data) < 100:
            break

        page += 1

    return repositories


def get_user_prs_from_repo(repo, username):
    prs = []
    page = 1

    while True:
        url = f"{GITHUB_API}/repos/{repo}/pulls"

        params = {
            "state": "all",
            "per_page": 100,
            "page": page,
        }

        data = github_request(url, params)

        if not data:
            break

        for pr in data:
            author = pr.get("user", {}).get("login")

            if author and author.lower() == username.lower():
                prs.append(pr)

        if len(data) < 100:
            break

        page += 1

    return prs


def get_all_prs(org, username):
    prs = []

    repositories = get_all_repositories(org)

    print(
        f"\nChecking {len(repositories)} repositories for PRs by @{username}..."
    )

    for index, repository in enumerate(
        repositories,
        start=1,
    ):
        repo = repository["full_name"]

        print(
            f"[{index}/{len(repositories)}] Checking {repo}..."
        )

        try:
            repo_prs = get_user_prs_from_repo(
                repo,
                username,
            )

            if repo_prs:
                prs.extend(repo_prs)

                print(
                    f"  Found {len(repo_prs)} PR(s) by @{username}"
                )

        except requests.HTTPError as e:
            print(
                f"  ERROR processing repository {repo}: {e}"
            )

    return prs


def get_pr_details(repo, pr_number):
    url = f"{GITHUB_API}/repos/{repo}/pulls/{pr_number}"

    return github_request(url)


def get_pr_commits(repo, pr_number):
    commits = []
    page = 1

    while True:
        url = (
            f"{GITHUB_API}/repos/"
            f"{repo}/pulls/{pr_number}/commits"
        )

        params = {
            "per_page": 100,
            "page": page,
        }

        data = github_request(url, params)

        if not data:
            break

        commits.extend(data)

        if len(data) < 100:
            break

        page += 1

    return commits


def calculate_duration(commits):
    if not commits:
        return None, None, None

    commit_dates = []

    for commit in commits:
        commit_data = commit.get("commit", {})
        author = commit_data.get("author")

        if author and author.get("date"):
            commit_dates.append(
                parse_datetime(author["date"])
            )

    if not commit_dates:
        return None, None, None

    first_commit = min(commit_dates)
    last_commit = max(commit_dates)

    duration = last_commit - first_commit
    days = duration.total_seconds() / 86400

    return first_commit, last_commit, days


def create_excel(rows, output_file):
    wb = Workbook()

    ws = wb.active
    ws.title = "My Pull Requests"

    headers = [
        "PR Number",
        "Repository",
        "Title",
        "Status",
        "First Commit",
        "Last Commit",
        "Days Taken",
        "Created At",
        "Closed At",
        "Merged At",
        "Author",
        "PR URL",
    ]

    ws.append(headers)

    for cell in ws[1]:
        cell.font = Font(bold=True)

        cell.fill = PatternFill(
            start_color="D9EAF7",
            end_color="D9EAF7",
            fill_type="solid",
        )

        cell.alignment = Alignment(
            horizontal="center"
        )

    for row in rows:
        ws.append([
            row["pr_number"],
            row["repository"],
            row["title"],
            row["status"],
            row["first_commit"],
            row["last_commit"],
            row["days_taken"],
            row["created_at"],
            row["closed_at"],
            row["merged_at"],
            row["author"],
            row["url"],
        ])

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(
                vertical="top"
            )

    date_columns = [5, 6, 8, 9, 10]

    for col in date_columns:
        for cell in ws.iter_cols(
            min_col=col,
            max_col=col,
            min_row=2,
        ):
            for c in cell:
                if c.value:
                    c.number_format = "yyyy-mm-dd hh:mm"

    for cell in ws["G"][1:]:
        if cell.value is not None:
            cell.number_format = "0.00"

    for cell in ws["L"][1:]:
        if cell.value:
            cell.hyperlink = cell.value
            cell.style = "Hyperlink"

    ws.freeze_panes = "A2"

    for column in ws.columns:
        max_length = 0

        column_letter = get_column_letter(
            column[0].column
        )

        for cell in column:
            try:
                length = len(str(cell.value))
                max_length = max(
                    max_length,
                    length,
                )
            except Exception:
                pass

        ws.column_dimensions[
            column_letter
        ].width = min(
            max_length + 2,
            50,
        )

    ws.auto_filter.ref = ws.dimensions

    wb.save(output_file)


def main():
    if not GITHUB_TOKEN:
        print(
            "ERROR: GITHUB_TOKEN is missing from .env"
        )
        sys.exit(1)

    if not GITHUB_ORG:
        print(
            "ERROR: GITHUB_ORG is missing from .env"
        )
        sys.exit(1)

    if not GITHUB_USERNAME:
        print(
            "ERROR: GITHUB_USERNAME is missing from .env"
        )
        sys.exit(1)

    output_file = "github_my_prs.xlsx"

    prs = get_all_prs(
        GITHUB_ORG,
        GITHUB_USERNAME,
    )

    print(
        f"\nTotal PRs found: {len(prs)}"
    )

    rows = []

    for index, pr in enumerate(
        prs,
        start=1,
    ):
        repo = pr["base"]["repo"]["full_name"]
        pr_number = pr["number"]

        print(
            f"[{index}/{len(prs)}] "
            f"{repo} #{pr_number}"
        )

        try:
            details = get_pr_details(
                repo,
                pr_number,
            )

            commits = get_pr_commits(
                repo,
                pr_number,
            )

            first_commit, last_commit, days = (
                calculate_duration(commits)
            )

            if details["state"] == "open":
                status = "Open"
            elif details.get("merged_at"):
                status = "Merged"
            else:
                status = "Closed"

            rows.append({
                "pr_number": pr_number,
                "repository": repo,
                "title": details.get(
                    "title",
                    "",
                ),
                "status": status,
                "first_commit": first_commit,
                "last_commit": last_commit,
                "days_taken": days,
                "created_at": parse_datetime(
                    details.get("created_at")
                ),
                "closed_at": parse_datetime(
                    details.get("closed_at")
                ),
                "merged_at": parse_datetime(
                    details.get("merged_at")
                ),
                "author": details.get(
                    "user",
                    {},
                ).get(
                    "login",
                    "",
                ),
                "url": details.get(
                    "html_url",
                    "",
                ),
            })

        except requests.HTTPError as e:
            print(
                f"  ERROR processing "
                f"{repo} #{pr_number}: {e}"
            )

    rows.sort(
        key=lambda x: (
            x["days_taken"]
            if x["days_taken"] is not None
            else -1
        ),
        reverse=True,
    )

    create_excel(
        rows,
        output_file,
    )

    print(
        "\n--------------------------------"
    )
    print("Done!")
    print(
        f"PRs exported: {len(rows)}"
    )
    print(
        f"Excel file: {output_file}"
    )
    print(
        "--------------------------------"
    )


if __name__ == "__main__":
    main()